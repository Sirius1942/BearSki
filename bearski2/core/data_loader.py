"""数据表加载器 - 支持动态函数 + JSON"""
import openpyxl
import json
from datetime import datetime
import random
import string

class DataLoader:
    """加载测试数据，支持函数和多种格式"""
    
    FUNCTIONS = {
        'now': lambda: datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'today': lambda: datetime.now().strftime('%Y-%m-%d'),
        'timestamp': lambda: int(datetime.now().timestamp()),
        'random_str': lambda n=8: ''.join(random.choices(string.ascii_letters, k=int(n))),
        'random_int': lambda min=1, max=100: random.randint(int(min), int(max)),
        'random_email': lambda: f"test_{random.randint(1000,9999)}@example.com",
    }
    
    def load(self, data_path: str):
        """统一加载接口，自动识别格式"""
        if '::' in data_path:
            # Excel 格式: testdata.xlsx::sheet::row_id
            file, sheet, row_id = data_path.split('::')
            if file.endswith('.xlsx'):
                return self.load_excel(file, sheet, row_id)
            elif file.endswith('.json'):
                return self.load_json(file, row_id)
        return {}
    
    def load_excel(self, file_path: str, sheet: str, row_id: str):
        """加载 Excel 数据"""
        wb = openpyxl.load_workbook(file_path)
        ws = wb[sheet]
        
        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == row_id:
                data = dict(zip(headers, [cell.value for cell in row]))
                return self._process_functions(data)
        return {}
    
    def load_json(self, file_path: str, data_id: str):
        """加载 JSON 数据"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data_list = json.load(f)
        
        for item in data_list:
            if item.get('DataID') == data_id:
                return self._process_functions(item)
        return {}
    
    def _process_functions(self, data: dict) -> dict:
        """处理数据中的函数 - 运行时动态计算"""
        result = {}
        for key, value in data.items():
            if isinstance(value, str) and value.startswith('${') and value.endswith('}'):
                func_expr = value[2:-1]
                result[key] = self._eval_function(func_expr)
            else:
                result[key] = value
        return result
    
    def _eval_function(self, expr: str):
        """执行函数表达式"""
        if '(' in expr:
            func_name = expr[:expr.index('(')]
            args_str = expr[expr.index('(')+1:expr.index(')')]
            args = [a.strip() for a in args_str.split(',')] if args_str else []
            return self.FUNCTIONS[func_name](*args)
        return self.FUNCTIONS[expr]()
